import unittest
from openpyxl import load_workbook
from settings import FILE_LOCATION, EXCEL_FILE
from excel_functions import *
from helper_func import convert_letter_to_index, check_if_already_exists

class Test_get_all_unique_ordernumbers(unittest.TestCase):
    
    def test_correct_letter_to_index(self):
        index = convert_letter_to_index("g")
        self.assertEqual(index, 6)

    def test_more_letters_to_index(self):
        index = convert_letter_to_index("p")
        self.assertEqual(index, 15)

    def test_get_all_unique_ordernumbers(self):
        wb = load_workbook(filename=f"{FILE_LOCATION}/{EXCEL_FILE}")
        extracted_order_nums = get_all_saved_unique_ordernumbers(wb, "Test historikk", "g")
        orders = {2246033,2246032,2246009,2245984,2241747}
        self.assertEqual(extracted_order_nums,orders)

class test_check_if_order_number_is_used(unittest.TestCase):

    def test_ord_num_exists_short_list(self):
        wb = load_workbook(filename=f"{FILE_LOCATION}/{EXCEL_FILE}")
        extracted_order_nums = get_all_saved_unique_ordernumbers(wb, "Test historikk", "g")
        exists = check_if_already_exists(extracted_order_nums, 2246033)
        self.assertEqual(exists, True)

    def test_not_ord_num_exists_short_list(self):
        wb = load_workbook(filename=f"{FILE_LOCATION}/{EXCEL_FILE}")
        extracted_order_nums = get_all_saved_unique_ordernumbers(wb, "Test historikk", "g")
        exists = check_if_already_exists(extracted_order_nums, 2246030)
        self.assertEqual(exists, False)

    def test_ord_num_exists_long_list(self):
        wb = load_workbook(filename=f"{FILE_LOCATION}/{EXCEL_FILE}")
        extracted_order_nums = get_all_saved_unique_ordernumbers(wb, "Test historikk lang", "g")
        exists = check_if_already_exists(extracted_order_nums, 2252742)  #2252742 should be the last number in list.
        self.assertEqual(exists, True)

class test_product_number_exists_in_database(unittest.TestCase):

    def test_product_number_exist(self):
        wb = load_workbook(filename=f"{FILE_LOCATION}/{EXCEL_FILE}")
        result = check_product_number_in_database(wb, "Test Database",0,21112)
        self.assertEqual(result, True)

    def test_product_number_not_exist(self):
        wb = load_workbook(filename=f"{FILE_LOCATION}/{EXCEL_FILE}")
        result = check_product_number_in_database(wb,"Test Database",0,21212)
        self.assertEqual(result, False)

class test_product_amounts_from_database(unittest.TestCase):
    def test_find_product_amount(self):
        wb = load_workbook(filename=f"{FILE_LOCATION}/{EXCEL_FILE}")
        value = find_product_amount(wb, "Test Database", 10001)
        self.assertEqual(value, 1893)

    def test_function_subtract(self):
        wb = load_workbook(filename=f"{FILE_LOCATION}/{EXCEL_FILE}")
        old_value = find_product_amount(wb, "Test Database", 10001)
        new_value = subtract_product_amount(wb, "Test Database", 10001, 5)
        self.assertEqual(old_value-5, new_value)

class test_add_to_history(unittest.TestCase):

    def test_add_to_bottom(self):
        wb = load_workbook(filename=f"{FILE_LOCATION}/{EXCEL_FILE}")
        historikk = wb["Test historikk"]
        bottom_row_usable = find_bottom_usable_row(wb, "Test historikk")
        sim_order = [2246101,[10002,"Bokashistrø, 2 kg",4]]
        add_to_bottom_history(wb, "Test historikk",sim_order[0],sim_order[1],5)
        self.assertEqual(bottom_row_usable, historikk.max_row)
        # TODO, figure out how to properly test bottom row
        
